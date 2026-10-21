"""测量报告：把测量记录填充到 xlsx 报告模板，再转 PDF。

报告模板为 runtime/pdf/abc.xlsx（用户提供），工作表 "report"，版式：
    B2:E2      标题：称重测量报告
    B3/C3      操作者 | D3/E3 时间
    B4/C4      车型   | D4/E4 车号
    B5         进车方向（值填 C5，C5:E5 合并）
    B6:E6      表头：轮号 | 轮重(kg) | 偏差(%) | 轴重(kg)
    B7:B24     轮号 1-18；C 列轮重(kg)；D 列偏差(%)；
               E 列轴重(kg)（E7:E8、E9:E10 … 两行合并为一轴）
    B25        整车重(kg)（值填 D25，D25:E25 合并）
    B26        不平衡检测(kg)（值填 D26，D26:E26 合并；= 最大轮重-最小轮重）

模板缺失时用 openpyxl 按同一版式生成（兜底，版式与 abc.xlsx 一致）。
数据库存储结构与模板 1:1 对应（见 dbstore），轮/轴/整车重/不平衡
均以 kg 存库，导出时直接填入，不再做结构转换与单位换算。
"""
from __future__ import annotations

import logging
from pathlib import Path
from typing import Any

from measure import KN_TO_KG  # noqa: F401  （兼容旧引用；换算系数定义在 measure）

logger = logging.getLogger(__name__)

TEMPLATE_SHEET = "report"


def ensure_measure_template(template_path: str | Path) -> Path:
    """模板不存在时按 abc.xlsx 版式生成兜底模板。返回模板路径。"""
    template_path = Path(template_path)
    if template_path.exists():
        return template_path
    template_path.parent.mkdir(parents=True, exist_ok=True)

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, Side

    wb = Workbook()
    ws = wb.active
    ws.title = TEMPLATE_SHEET
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    title_font = Font(size=16, bold=True)
    head_font = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center")
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.merge_cells("B2:E2")
    c = ws["B2"]
    c.value = "称重测量报告"
    c.font = title_font
    c.alignment = center

    for coord, label in [("B3", "操作者"), ("D3", "时间"),
                         ("B4", "车型"), ("D4", "车号"),
                         ("B5", "进车方向")]:
        ws[coord].value = label
        ws[coord].font = head_font
    ws.merge_cells("C5:E5")

    for coord, label in [("B6", "轮号"), ("C6", "轮重(kg)"),
                         ("D6", "偏差(%)"), ("E6", "轴重(kg)")]:
        ws[coord].value = label
        ws[coord].font = head_font
        ws[coord].alignment = center

    for i in range(18):  # B7:B24 轮号
        r = 7 + i
        ws[f"B{r}"].value = i + 1
        ws[f"B{r}"].alignment = center
    for k in range(9):  # 轴重两行合并
        ws.merge_cells(f"E{7 + 2 * k}:E{8 + 2 * k}")

    ws["B25"].value = "整车重(kg)"
    ws["B25"].font = head_font
    ws.merge_cells("B25:C25")
    ws.merge_cells("D25:E25")
    ws["B26"].value = "不平衡检测(kg)"
    ws["B26"].font = head_font
    ws.merge_cells("B26:C26")
    ws.merge_cells("D26:E26")

    for row in ws["B2:E26"]:
        for cell in row:
            cell.border = border

    for col, width in zip("ABCDE", [2, 14, 16, 14, 16]):
        ws.column_dimensions[col].width = width

    wb.save(template_path)
    logger.info("生成测量报告模板: %s", template_path)
    return template_path


def _num(v: Any) -> float | str:
    """存储值 -> 单元格值：数值原样，空串/None 保持空白。"""
    return round(float(v), 2) if isinstance(v, (int, float)) else ""


def build_measure_cells(record: dict[str, Any], operator: str = "") -> dict[str, Any]:
    """记录 -> {单元格引用: 值}，供 PdfReport.set_cells 使用（abc.xlsx 版式）。

    记录的存储结构与模板 1:1 对应（轮数据/轴数据/整车重/不平衡均为
    kg），本函数只做单元格定位，不做结构转换与单位换算。操作者优先
    用记录自带的（测量时落库），为空时用调用方传入的当前用户兜底。
    """
    cells: dict[str, Any] = {
        "C3": record.get("operator") or operator,
        "E3": record.get("measure_time", ""),
        "C4": record.get("model", ""),
        "E4": record.get("car_no", ""),
        "C5": record.get("direction", ""),  # 进车方向
    }
    for entry in record.get("wheels") or []:
        r = 6 + int(entry["轮号"])
        cells[f"C{r}"] = _num(entry.get("轮重(kg)"))
        cells[f"D{r}"] = _num(entry.get("偏差(%)"))
    for entry in record.get("axles") or []:
        v = entry.get("轴重(kg)")
        cells[f"E{7 + 2 * (int(entry['轴号']) - 1)}"] = _num(v)

    total = record.get("total_kg")
    cells["D25"] = round(float(total), 1) if isinstance(total, (int, float)) else ""
    balance = record.get("balance_kg")
    cells["D26"] = round(float(balance), 1) if isinstance(balance, (int, float)) else ""
    return cells


def export_measure_pdf(
    record: dict[str, Any],
    template_path: str | Path,
    output_dir: str | Path,
    filename: str | None = None,
    operator: str = "",
) -> Path:
    """导出单条测量记录的 PDF 报告，返回 PDF 路径。"""
    from pdf import PdfReport

    template_path = ensure_measure_template(template_path)
    report = PdfReport(template_path, output_dir)
    report.load_template(TEMPLATE_SHEET)
    report.set_cells(build_measure_cells(record, operator=operator))
    if filename is None:
        rid = record.get("id", "0")
        car = str(record.get("car_no", "")).strip() or "nocar"
        filename = f"称重报告_{car}_{rid}.pdf"
    return report.export_pdf(filename=filename, title="称重测量报告")
